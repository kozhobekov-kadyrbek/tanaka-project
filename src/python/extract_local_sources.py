"""Extract publication-safe derived tables from the locally held Tanaka workbooks.

The workbooks are restricted source artefacts and are intentionally ignored by git.
This utility reads their cached calculation results and writes only numerical derived
tables, so a reviewer of the public repository never needs the originals.
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


REGIONS = ["北海道", "東北", "東京", "北陸", "中部", "関西", "中国", "四国", "九州", "沖縄"]


def number(value: object) -> float:
    if value is None:
        raise ValueError("A required cached workbook value is blank")
    return float(value)


def scenario_labels(scenario: int) -> tuple[str, str, str]:
    ef = "carbon_only" if scenario <= 12 else "carbon_and_food"
    thermal = ["2016_mix", "2030_mix", "lng_only", "zero_thermal"][(scenario - 1) % 12 // 3]
    nuclear = ["currently_operating", "all_candidate_units", "zero_nuclear"][(scenario - 1) % 3]
    return ef, thermal, nuclear


def write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def extract_final(final_workbook: Path, output_root: Path) -> None:
    wb = load_workbook(final_workbook, data_only=True, read_only=True)
    try:
        gross = wb["再エネ目標量"]
        shortfall = wb["必要再エネ量"]
        reported: list[dict] = []
        for scenario in range(1, 25):
            ef, thermal, nuclear = scenario_labels(scenario)
            gross_col = scenario + 2
            # Every new 6-scenario block in the shortfall sheet contains a label column.
            shortfall_col = scenario + 2 + (scenario >= 7) + (scenario >= 13) + (scenario >= 19)
            for row in range(3, 13):
                region = gross.cell(row, 2).value
                if region not in REGIONS:
                    raise ValueError(f"Unexpected region at 再エネ目標量!B{row}: {region!r}")
                g = number(gross.cell(row, gross_col).value) / 1_000_000
                net = number(shortfall.cell(row + 1, shortfall_col).value) / 1_000_000
                reported.append({
                    "scenario": scenario, "EF_case": ef, "thermal_case": thermal,
                    "nuclear_case": nuclear, "region": region,
                    "gross_RE_requirement_TWh": f"{g:.6f}",
                    "signed_net_RE_shortfall_TWh": f"{net:.6f}",
                    "positive_additional_RE_TWh": f"{max(0, net):.6f}",
                    "source_cell_gross": f"再エネ目標量!{gross.cell(row, gross_col).coordinate}",
                    "source_cell_net": f"必要再エネ量!{shortfall.cell(row + 1, shortfall_col).coordinate}",
                })
        write_csv(
            output_root / "data/derived/canonical_final_workbook/as_reported_240.csv",
            list(reported[0]), reported,
        )

        cap_sheet = wb["CO2削減エリア"]
        caps = []
        for row in range(21, 31):
            region = cap_sheet.cell(row, 10).value
            caps.append({
                "region": region,
                "carbon_only_cap_mt": f"{number(cap_sheet.cell(row, 11).value) / 1_000_000:.12f}",
                "carbon_and_food_cap_mt": f"{number(cap_sheet.cell(row, 12).value) / 1_000_000:.12f}",
                "baseline_co2_mt": f"{number(cap_sheet.cell(row, 13).value) / 1_000_000:.12f}",
                "source_row": row,
            })
        write_csv(output_root / "data/derived/canonical_final_workbook/regional_caps_from_final.csv", list(caps[0]), caps)
    finally:
        wb.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--final-workbook", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args()
    extract_final(args.final_workbook, args.output_root)


if __name__ == "__main__":
    main()
